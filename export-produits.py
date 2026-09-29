# -*- coding: utf-8 -*-
"""
Exporte la base produits (déduupliquée, comme l'import Firestore) vers Excel :
C:/Users/usre/Desktop/projet de prospections/Produits_TRADE_LINK.xlsx

Même clé de dédoublonnage que import-produits.mjs (accents/pluriel ignorés).
Usage : python export-produits.py
"""
import json
import re
import unicodedata
from hashlib import md5
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ERP = r"C:\Users\usre\Desktop\projet de prospections"
BASE = ERP + r"\..\Facture TRADE LINK INTERNATIONALE\base_produits.json"
OUT = ERP + r"\Produits_TRADE_LINK.xlsx"


def dedup_key(name):
    s = unicodedata.normalize("NFD", name or "")
    s = "".join(c for c in s if not unicodedata.combining(c)).upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    s = " ".join(w[:-1] if len(w) > 3 and w.endswith("S") else w for w in s.split(" "))
    return s


def id_for(key):
    return "imp" + md5(key.encode()).hexdigest()[:12]


def main():
    base = json.load(open(BASE, encoding="utf-8"))
    merged = {}
    services = 0
    fusionnes = 0
    for key, p in base.items():
        k = dedup_key(p.get("designation") or key)
        if not k:
            continue
        if p.get("type") == "service":
            services += 1
            continue
        pu = p.get("pu_dernier") or p.get("pu_moyen")
        m = merged.get(k)
        if m is None:
            merged[k] = {
                "nom": p.get("designation"), "desc": p.get("description") or "",
                "prix": pu, "pmin": p.get("pu_min"), "pmax": p.get("pu_max"),
                "poids": p.get("poids_unitaire_kg"),
                "date": p.get("derniere_vue") or "",
                "nbf": len(p.get("factures") or []),
            }
        else:
            fusionnes += 1
            if pu and (p.get("derniere_vue") or "") >= m["date"]:
                m["prix"], m["date"] = pu, p.get("derniere_vue") or m["date"]
            m["pmin"] = min([x for x in (m["pmin"], p.get("pu_min")) if x is not None], default=None)
            m["pmax"] = max([x for x in (m["pmax"], p.get("pu_max")) if x is not None], default=None)
            if not m["poids"] and p.get("poids_unitaire_kg"):
                m["poids"] = p["poids_unitaire_kg"]
            if not m["desc"] and p.get("description"):
                m["desc"] = p["description"]
            m["nbf"] += len(p.get("factures") or [])

    wb = Workbook()
    ws = wb.active
    ws.title = "Produits"
    heads = ["ID Firestore", "Nom produit", "Description", "Prix de vente (MAD)",
             "Prix min facturé", "Prix max facturé", "Poids unitaire (kg)",
             "Nb factures", "Dernière facture"]
    ws.append(heads)
    fill = PatternFill("solid", fgColor="1F4E78")
    for c in range(1, len(heads) + 1):
        cell = ws.cell(1, c)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(vertical="center")
    for k in sorted(merged):
        p = merged[k]
        ws.append([id_for(k), p["nom"], p["desc"][:200], p["prix"],
                   p["pmin"], p["pmax"], p["poids"], p["nbf"], p["date"]])
        for c in (4, 5, 6, 7):
            ws.cell(ws.max_row, c).number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHI", (16, 44, 40, 15, 13, 13, 14, 10, 13)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{ws.max_row}"

    wb.save(OUT)
    print(f"✔ {len(merged)} produits uniques exportés ({fusionnes} doublons fusionnés, "
          f"{services} services exclus)\n→ {OUT}")


if __name__ == "__main__":
    main()
